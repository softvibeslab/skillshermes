#!/usr/bin/env python3
"""Lightweight QA for python-pptx/PPTX decks.

Checks:
- slide count
- text boxes outside 16:9 widescreen bounds
- long text in very short boxes (rough overflow heuristic)
- tiny text boxes that are likely accidental

Usage:
  python scripts/qa_pptx_textboxes.py /path/to/deck.pptx
"""
from __future__ import annotations
import sys
from pptx import Presentation
from pptx.util import Inches

SLIDE_W = 13.333
SLIDE_H = 7.5


def main(path: str) -> int:
    prs = Presentation(path)
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not hasattr(sh, "text") or not sh.text.strip():
                continue
            x = sh.left / Inches(1)
            y = sh.top / Inches(1)
            w = sh.width / Inches(1)
            h = sh.height / Inches(1)
            txt = sh.text.strip().replace("\n", " ")
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                issues.append((si, "out of bounds", txt[:80], round(x, 2), round(y, 2), round(w, 2), round(h, 2)))
            if len(txt) > 120 and h < 0.6:
                issues.append((si, "possible text overflow", txt[:80], round(x, 2), round(y, 2), round(w, 2), round(h, 2)))
            if w < 0.25 or h < 0.1:
                issues.append((si, "tiny text box", txt[:80], round(x, 2), round(y, 2), round(w, 2), round(h, 2)))
    print(f"slides {len(prs.slides)} issues {len(issues)}")
    for issue in issues:
        print(issue)
    return 1 if issues else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("usage: qa_pptx_textboxes.py /path/to/deck.pptx", file=sys.stderr)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
